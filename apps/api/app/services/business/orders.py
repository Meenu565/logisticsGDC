import csv
from datetime import datetime
from decimal import Decimal
from io import BytesIO, StringIO
from uuid import UUID

from openpyxl import load_workbook

from app.models import Order, OrderImport, Stop
from app.models.enums import ImportStatus, OrderStatus, StopKind
from app.repositories.orders import OrderRepository
from app.services.business.events import DispatchEventService


class OrderService:
    def __init__(self, orders: OrderRepository, events: DispatchEventService | None = None) -> None:
        self.orders = orders
        self.events = events

    def list_active_orders(self) -> list[Order]:
        return list(self.orders.list_active())

    def get_orders_for_planning(self, order_ids: list[UUID] | None = None) -> list[Order]:
        if order_ids:
            return list(self.orders.list_by_ids(order_ids))
        return self.list_active_orders()

    def get_order(self, order_id: UUID) -> Order:
        order = self.orders.get(order_id)
        if order is None:
            raise ValueError(f"Order not found: {order_id}")
        return order

    def load_order_range(self, start: int, end: int) -> list[Order]:
        if start > end:
            raise ValueError("Range start must be less than or equal to range end.")
        return list(self.orders.list_external_numeric_range(start, end))

    def update_order(
        self,
        order_id: UUID,
        updates: dict,
        session_id: UUID | None = None,
    ) -> Order:
        allowed_fields = {
            "priority",
            "customer_name",
            "required_weight_kg",
            "required_volume_m3",
            "time_window_start",
            "time_window_end",
            "notes",
            "status",
        }
        order = self.get_order(order_id)
        for field, value in updates.items():
            if field not in allowed_fields:
                raise ValueError(f"Unsupported order update field: {field}")
            setattr(order, field, value)
        order.updated_at = datetime.utcnow()
        self.orders.add(order)
        if self.events:
            self.events.record(
                event_type="order_updated",
                title=f"Order updated: {order.external_order_id or order.id}",
                details={"updates": updates},
                session_id=session_id,
                order_id=order.id,
            )
        return order

    def import_orders_csv(
        self,
        csv_text: str,
        file_name: str,
        session_id: UUID | None = None,
        mode: str = "append",
    ) -> dict:
        reader = csv.DictReader(StringIO(csv_text))
        return self.import_order_rows(list(reader), file_name=file_name, source_type="csv", session_id=session_id, mode=mode)

    def import_orders_excel(
        self,
        content: bytes,
        file_name: str,
        session_id: UUID | None = None,
        mode: str = "append",
    ) -> dict:
        workbook = load_workbook(filename=BytesIO(content))
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return {"file_name": file_name, "imported_rows": 0, "mode": mode, "replaced_rows": 0}
        headers = [str(value).strip() for value in rows[0]]
        records = [
            {headers[index]: value for index, value in enumerate(row)}
            for row in rows[1:]
            if any(value is not None for value in row)
        ]
        return self.import_order_rows(records, file_name=file_name, source_type="excel", session_id=session_id, mode=mode)

    def import_order_rows(
        self,
        rows: list[dict],
        file_name: str,
        source_type: str,
        session_id: UUID | None = None,
        mode: str = "append",
    ) -> dict:
        order_import = self.orders.add_import(
            OrderImport(
                file_name=file_name,
                source_type=source_type,
                status=ImportStatus.PENDING,
                total_rows=len(rows),
            )
        )
        imported_rows = 0
        replaced_rows = self.orders.cancel_active() if mode == "replace" else 0
        try:
            for row in rows:
                order = self._create_order_from_row(row, order_import.id)
                self._create_optional_stops(order, row)
                imported_rows += 1
            self.orders.complete_import(order_import, imported_rows)
        except Exception as exc:
            self.orders.complete_import(order_import, imported_rows, ImportStatus.FAILED, str(exc))
            raise
        if self.events:
            self.events.record(
                event_type="orders_imported",
                title=f"Orders imported from {file_name}",
                details={
                    "file_name": file_name,
                    "source_type": source_type,
                    "imported_rows": imported_rows,
                    "mode": mode,
                    "replaced_rows": replaced_rows,
                },
                session_id=session_id,
            )
        return {
            "file_name": file_name,
            "source_type": source_type,
            "imported_rows": imported_rows,
            "mode": mode,
            "replaced_rows": replaced_rows,
        }

    def _create_order_from_row(self, row: dict, import_id: UUID) -> Order:
        delivery_type_code = str(row.get("delivery_type") or "curbside").strip().lower()
        delivery_type = self.orders.get_delivery_type_by_code(delivery_type_code)
        if delivery_type is None:
            raise ValueError(f"Unknown delivery type: {delivery_type_code}")
        order = Order(
            external_order_id=self._text(row.get("external_order_id")),
            import_id=import_id,
            delivery_type_id=delivery_type.id,
            customer_name=str(row["customer_name"]).strip(),
            priority=int(row.get("priority") or 0),
            required_weight_kg=Decimal(str(row.get("required_weight_kg") or 0)),
            required_volume_m3=Decimal(str(row.get("required_volume_m3") or 0)),
            notes=self._text(row.get("notes")),
        )
        return self.orders.add(order)

    def _create_optional_stops(self, order: Order, row: dict) -> None:
        for prefix, kind in [("pickup", StopKind.PICKUP), ("dropoff", StopKind.DROPOFF)]:
            address = self._text(row.get(f"{prefix}_address_line1"))
            city = self._text(row.get(f"{prefix}_city"))
            if not address or not city:
                continue
            self.orders.add_stop(
                Stop(
                    order_id=order.id,
                    kind=kind,
                    address_line1=address,
                    address_line2=self._text(row.get(f"{prefix}_address_line2")),
                    city=city,
                    region=self._text(row.get(f"{prefix}_region")),
                    postal_code=self._text(row.get(f"{prefix}_postal_code")),
                    country=self._text(row.get(f"{prefix}_country")) or "Canada",
                    latitude=self._decimal(row.get(f"{prefix}_latitude")),
                    longitude=self._decimal(row.get(f"{prefix}_longitude")),
                )
            )

    def _text(self, value: object) -> str | None:
        if value is None:
            return None
        text = str(value).strip()
        return text or None

    def _decimal(self, value: object) -> Decimal | None:
        text = self._text(value)
        return Decimal(text) if text is not None else None

    def cancel_order(self, order_id: UUID, session_id: UUID | None = None, reason: str | None = None) -> Order:
        order = self.orders.get(order_id)
        if order is None:
            raise ValueError(f"Order not found: {order_id}")
        order.status = OrderStatus.CANCELLED
        order.cancelled_at = datetime.utcnow()
        self.orders.add(order)
        if self.events:
            self.events.record(
                event_type="order_cancelled",
                title=f"Order cancelled: {order.external_order_id or order.id}",
                details={"reason": reason},
                session_id=session_id,
                order_id=order.id,
            )
        return order
